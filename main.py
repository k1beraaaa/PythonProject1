"""Демо-бот записи со свободными слотами.

Клиент: «Записаться» -> имя -> услуга -> дата (кнопки) -> время (кнопки со свободными слотами).
Владельцу приходит заявка в Telegram, строка пишется в Excel и (по желанию) в Google Таблицу.
Занятые слоты хранятся в bookings.json.
"""
import asyncio
import json
import logging
import os
from datetime import date, datetime, time, timedelta
from pathlib import Path

from aiogram import Bot, Dispatcher, F
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
)
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

load_dotenv()
load_dotenv("settings.env")

BOT_TOKEN = os.environ["BOT_TOKEN"]
OWNER_CHAT_ID = int(os.getenv("OWNER_CHAT_ID", "0"))
SERVICES = [s.strip() for s in os.getenv("SERVICES", "Стрижка,Маникюр,Консультация").split(",")]
EXCEL_PATH = Path(os.getenv("EXCEL_PATH", "zayavki.xlsx"))
GSHEET_ID = os.getenv("GSHEET_ID")
GOOGLE_CREDS = os.getenv("GOOGLE_CREDS", "service_account.json")

# --- расписание ---
SLOTS_FILE = Path(os.getenv("SLOTS_FILE", "bookings.json"))
WORK_START = int(os.getenv("WORK_START", "10"))        # начало работы, час
WORK_END = int(os.getenv("WORK_END", "19"))            # конец работы, час
SLOT_MINUTES = int(os.getenv("SLOT_MINUTES", "60"))    # длина слота
DAYS_AHEAD = int(os.getenv("DAYS_AHEAD", "7"))         # на сколько дней вперёд
DAYS_OFF = {int(x) for x in os.getenv("DAYS_OFF", "6").split(",") if x.strip()}  # 0=Пн ... 6=Вс

WEEKDAYS = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]
HEADERS = ["Дата заявки", "Имя", "Услуга", "Дата и время записи", "Telegram", "Chat ID"]
LOCK = asyncio.Lock()

logging.basicConfig(level=logging.INFO)
dp = Dispatcher()


class Booking(StatesGroup):
    name = State()
    service = State()
    slot = State()


# ---------- слоты ----------
def load_booked() -> set:
    if SLOTS_FILE.exists():
        return set(json.loads(SLOTS_FILE.read_text(encoding="utf-8")))
    return set()


def add_booked(key: str) -> None:
    booked = load_booked()
    booked.add(key)
    SLOTS_FILE.write_text(
        json.dumps(sorted(booked), ensure_ascii=False, indent=1), encoding="utf-8"
    )


def free_slots(day: date) -> list:
    booked = load_booked()
    now = datetime.now()
    result = []
    t = datetime.combine(day, time(WORK_START))
    end = datetime.combine(day, time(WORK_END))
    while t + timedelta(minutes=SLOT_MINUTES) <= end:
        key = t.strftime("%Y-%m-%d %H:%M")
        if t > now + timedelta(minutes=30) and key not in booked:
            result.append(key)
        t += timedelta(minutes=SLOT_MINUTES)
    return result


def days_with_free() -> list:
    days = []
    for i in range(DAYS_AHEAD):
        day = date.today() + timedelta(days=i)
        if day.weekday() not in DAYS_OFF and free_slots(day):
            days.append(day)
    return days


def human(key: str) -> str:
    dt = datetime.strptime(key, "%Y-%m-%d %H:%M")
    return f"{dt:%d.%m.%Y %H:%M} ({WEEKDAYS[dt.weekday()]})"


def chunk(buttons: list, size: int) -> list:
    return [buttons[i:i + size] for i in range(0, len(buttons), size)]


def dates_kb(days: list) -> InlineKeyboardMarkup:
    btns = [
        InlineKeyboardButton(
            text=f"{WEEKDAYS[d.weekday()]} {d:%d.%m}", callback_data=f"d:{d.isoformat()}"
        )
        for d in days
    ]
    return InlineKeyboardMarkup(inline_keyboard=chunk(btns, 3))


def times_kb(slots: list) -> InlineKeyboardMarkup:
    btns = [
        InlineKeyboardButton(text=s[-5:], callback_data=f"t:{s}") for s in slots
    ]
    rows = chunk(btns, 3)
    rows.append([InlineKeyboardButton(text="← Другая дата", callback_data="back")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def start_kb() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text="Записаться")]], resize_keyboard=True
    )


def services_kb() -> ReplyKeyboardMarkup:
    rows = [[KeyboardButton(text=s)] for s in SERVICES]
    return ReplyKeyboardMarkup(keyboard=rows, resize_keyboard=True, one_time_keyboard=True)


# ---------- таблицы ----------
COL_WIDTHS = (18, 20, 20, 28, 20, 14)
HEADER_COLOR = "264773"   # тёмно-синий
ZEBRA_COLOR = "EEF3FA"    # светло-голубой


def save_to_excel(row: list) -> None:
    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    if EXCEL_PATH.exists():
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Заявки"
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF", size=11)
            c.fill = PatternFill("solid", start_color=HEADER_COLOR)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        ws.row_dimensions[1].height = 28
        for col, width in zip("ABCDEF", COL_WIDTHS):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

    ws.append(row)
    r = ws.max_row
    for c in ws[r]:
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = border
        if r % 2 == 1:
            c.fill = PatternFill("solid", start_color=ZEBRA_COLOR)
    ws.row_dimensions[r].height = 22
    ws.auto_filter.ref = f"A1:F{r}"
    wb.save(EXCEL_PATH)


def format_gsheet(ws) -> None:
    """Оформление Google Таблицы: делается один раз, когда лист пустой."""
    def rgb(hex_):
        return {
            "red": int(hex_[0:2], 16) / 255,
            "green": int(hex_[2:4], 16) / 255,
            "blue": int(hex_[4:6], 16) / 255,
        }

    requests = [
        {
            "addBanding": {
                "bandedRange": {
                    "range": {
                        "sheetId": ws.id,
                        "startRowIndex": 0,
                        "endRowIndex": 1000,
                        "startColumnIndex": 0,
                        "endColumnIndex": len(HEADERS),
                    },
                    "rowProperties": {
                        "headerColor": rgb(HEADER_COLOR),
                        "firstBandColor": rgb("FFFFFF"),
                        "secondBandColor": rgb(ZEBRA_COLOR),
                    },
                }
            }
        },
        {
            "updateSheetProperties": {
                "properties": {"sheetId": ws.id, "gridProperties": {"frozenRowCount": 1}},
                "fields": "gridProperties.frozenRowCount",
            }
        },
        {"setBasicFilter": {"filter": {"range": {
            "sheetId": ws.id,
            "startRowIndex": 0,
            "startColumnIndex": 0,
            "endColumnIndex": len(HEADERS),
        }}}},
    ]
    for i, w in enumerate(COL_WIDTHS):
        requests.append({
            "updateDimensionProperties": {
                "range": {"sheetId": ws.id, "dimension": "COLUMNS",
                          "startIndex": i, "endIndex": i + 1},
                "properties": {"pixelSize": int(w * 7.5)},
                "fields": "pixelSize",
            }
        })
    ws.spreadsheet.batch_update({"requests": requests})
    ws.format("A1:F1", {
        "textFormat": {"bold": True, "foregroundColor": rgb("FFFFFF")},
        "horizontalAlignment": "CENTER",
        "verticalAlignment": "MIDDLE",
    })


def save_to_gsheet(row: list) -> None:
    import gspread

    gc = gspread.service_account(filename=GOOGLE_CREDS)
    ws = gc.open_by_key(GSHEET_ID).sheet1
    if not ws.get_all_values():
        ws.append_row(HEADERS)
        try:
            format_gsheet(ws)
        except Exception:
            logging.exception("Не удалось оформить Google Таблицу")
    ws.append_row(row, value_input_option="USER_ENTERED")


async def save_row(row: list) -> None:
    await asyncio.to_thread(save_to_excel, row)
    if GSHEET_ID:
        try:
            await asyncio.to_thread(save_to_gsheet, row)
        except Exception:
            logging.exception("Не удалось записать в Google Таблицу")


# ---------- диалог ----------
@dp.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.clear()
    await message.answer(
        "Здравствуйте! Нажмите «Записаться», и я оформлю заявку за полминуты.",
        reply_markup=start_kb(),
    )


@dp.message(Command("myid"))
async def cmd_myid(message: Message):
    await message.answer(f"Ваш chat id: {message.chat.id}")


@dp.message(F.text.casefold() == "записаться")
async def begin(message: Message, state: FSMContext):
    await state.set_state(Booking.name)
    await message.answer("Как вас зовут?", reply_markup=ReplyKeyboardRemove())


@dp.message(Booking.name, F.text)
async def got_name(message: Message, state: FSMContext):
    await state.update_data(name=message.text.strip())
    await state.set_state(Booking.service)
    await message.answer("Какая услуга вас интересует?", reply_markup=services_kb())


@dp.message(Booking.service, F.text)
async def got_service(message: Message, state: FSMContext):
    await state.update_data(service=message.text.strip())
    await state.set_state(Booking.slot)
    days = days_with_free()
    if not days:
        await state.clear()
        await message.answer(
            "К сожалению, свободных мест сейчас нет. Попробуйте позже.",
            reply_markup=start_kb(),
        )
        return
    await message.answer("Выберите удобную дату:", reply_markup=dates_kb(days))


@dp.callback_query(Booking.slot, F.data.startswith("d:"))
async def pick_date(call: CallbackQuery):
    day = date.fromisoformat(call.data[2:])
    slots = free_slots(day)
    if not slots:
        days = days_with_free()
        await call.message.edit_text(
            "На этот день места закончились. Выберите другую дату:",
            reply_markup=dates_kb(days) if days else None,
        )
    else:
        await call.message.edit_text(
            f"Свободное время на {day:%d.%m} ({WEEKDAYS[day.weekday()]}):",
            reply_markup=times_kb(slots),
        )
    await call.answer()


@dp.callback_query(Booking.slot, F.data == "back")
async def back_to_dates(call: CallbackQuery):
    days = days_with_free()
    await call.message.edit_text("Выберите удобную дату:", reply_markup=dates_kb(days))
    await call.answer()


@dp.callback_query(Booking.slot, F.data.startswith("t:"))
async def pick_time(call: CallbackQuery, state: FSMContext, bot: Bot):
    key = call.data[2:]

    # проверка и бронь под замком, чтобы два клиента не заняли один слот
    async with LOCK:
        if key in load_booked():
            taken = True
        else:
            taken = False
            add_booked(key)

    if taken:
        day = datetime.strptime(key, "%Y-%m-%d %H:%M").date()
        slots = free_slots(day)
        await call.answer("Это время только что заняли, выберите другое", show_alert=True)
        await call.message.edit_text(
            f"Свободное время на {day:%d.%m}:" if slots else "На этот день места закончились.",
            reply_markup=times_kb(slots) if slots else dates_kb(days_with_free()),
        )
        return

    data = await state.get_data()
    await state.clear()

    user = call.from_user
    username = f"@{user.username}" if user.username else user.full_name
    now = datetime.now().strftime("%d.%m.%Y %H:%M")
    when = human(key)

    await save_row([now, data["name"], data["service"], when, username, call.message.chat.id])

    if OWNER_CHAT_ID:
        await bot.send_message(
            OWNER_CHAT_ID,
            "🔔 Новая заявка\n\n"
            f"👤 Имя: {data['name']}\n"
            f"💇 Услуга: {data['service']}\n"
            f"🕒 Время: {when}\n"
            f"✉️ Telegram: {username}\n"
            f"🗓 Получена: {now}",
        )

    await call.message.edit_text(f"Вы записаны на {when} ✅\nУслуга: {data['service']}")
    await call.message.answer(
        "Спасибо! Мы свяжемся с вами для подтверждения.", reply_markup=start_kb()
    )
    await call.answer()


async def main():
    bot = Bot(BOT_TOKEN)
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())